"""Inspect and convert manually labeled benchmark Excel files to JSONL.

The source workbooks use row fill colors as the human label.  This utility
keeps the source files untouched and emits one normalized record per row.
It deliberately does not use a source ``情感`` column as the gold label.
"""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


BLUE = {"FF00B0F0", "00FF00B0F0"}
YELLOW = {"FFFFFF00", "00FFFFFF00"}
NONE = {"", "00000000", "000000", "FFFFFFFF", "00FFFFFF"}

TEXT_ALIASES = {
    "标题": "title",
    "题目": "title",
    "正文": "text",
    "内容": "text",
    "封面ocr": "cover_ocr",
    "封面文字": "cover_ocr",
    "内容ocr": "content_ocr",
    "内容文字": "content_ocr",
}

LABEL_ALIASES = {
    "颜色": "label",
    "标签": "label",
    "标注": "label",
    "人工标注": "label",
    "gold": "label",
    "label": "label",
}

BRAND_ALIASES = {
    "百威": ["百威", "budweiser", "bud"],
    "哈尔滨": ["哈尔滨啤酒", "哈啤", "harbinbeer"],
    "科罗娜": ["corona", "科罗娜", "克罗那"],
    "雪津": ["雪津", "sedrin"],
    "青岛": ["青岛啤酒", "青啤", "崂山啤酒", "青岛纯生", "青岛1903", "青岛奥古特"],
    "雪花": ["雪花啤酒", "勇闯天涯", "superx", "snowbeer"],
    "乌苏": ["乌苏", "wusu", "红乌苏", "大乌苏", "夺命大乌苏"],
    "喜力": ["heineken", "喜力", "海尼根"],
    "RIO": ["锐澳", "RIO鸡尾酒", "锐澳鸡尾酒", "rio微醺", "rio强爽"],
    "乐堡": ["乐堡", "tuborg"],
}


def clean(value: Any) -> str:
    if value is None:
        return ""
    return str(value).replace("\ufeff", "").replace("\u200b", "").strip()


def normalize_header(value: Any) -> str:
    return re.sub(r"\s+", "", clean(value)).lower()


def cell_color(cell) -> str:
    color = cell.fill.fgColor
    value = color.rgb or ""
    return value.upper()


def label_from_color(row_cells) -> str | None:
    colors = [cell_color(cell) for cell in row_cells]
    if any(color in BLUE for color in colors):
        return "blue"
    if any(color in YELLOW for color in colors):
        return "yellow"
    if all(color in NONE for color in colors):
        return "none"
    return None


def label_from_value(value: Any) -> str | None:
    key = clean(value).lower().replace(" ", "")
    if key in {"blue", "蓝", "蓝色", "本品"}:
        return "blue"
    if key in {"yellow", "黄", "黄色", "竞品", "行业"}:
        return "yellow"
    if key in {"none", "不标", "无", "不标识", "中性"}:
        return "none"
    return None


def choose_columns(headers: list[str]) -> dict[str, str]:
    chosen: dict[str, str] = {}
    for header in headers:
        normalized = normalize_header(header)
        field = TEXT_ALIASES.get(normalized)
        if field and field not in chosen:
            chosen[field] = header
    return chosen


def choose_label_column(headers: list[str]) -> str | None:
    for header in headers:
        if LABEL_ALIASES.get(normalize_header(header)) == "label":
            return header
    return None


def infer_category(source_file: str, label: str) -> str:
    if label == "blue":
        return "own"
    if label == "yellow":
        return "industry" if source_file.startswith("行业") else "competitor"
    return "none"


def extract_brands(text: str) -> list[str]:
    normalized = clean(text).lower()
    return [
        brand
        for brand, aliases in BRAND_ALIASES.items()
        if any(alias.lower() in normalized for alias in aliases)
    ]


def iter_workbooks(input_dir: Path):
    for path in sorted(input_dir.rglob("*.xlsx")):
        if path.name.startswith("~$"):
            continue
        yield path


def inspect(input_dir: Path) -> int:
    total = 0
    for path in iter_workbooks(input_dir):
        workbook = load_workbook(path, read_only=False, data_only=True)
        print(f"{path.name}")
        for worksheet in workbook.worksheets:
            headers = [
                worksheet.cell(1, column).value
                for column in range(1, worksheet.max_column + 1)
            ]
            colors: dict[str, int] = {}
            for row in worksheet.iter_rows(min_row=2):
                color = label_from_color(row)
                colors[color or "unrecognized"] = colors.get(color or "unrecognized", 0) + 1
            print(
                f"  sheet={worksheet.title!r} rows={worksheet.max_row - 1} "
                f"columns={worksheet.max_column} colors={colors}"
            )
            print(f"  headers={headers}")
            total += max(0, worksheet.max_row - 1)
    print(f"total_rows={total}")
    return 0


def convert(input_dir: Path, output_path: Path) -> int:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    records: list[dict[str, Any]] = []
    skipped: list[str] = []
    for path in iter_workbooks(input_dir):
        workbook = load_workbook(path, read_only=False, data_only=True)
        for worksheet in workbook.worksheets:
            headers = [
                clean(worksheet.cell(1, column).value)
                for column in range(1, worksheet.max_column + 1)
            ]
            columns = choose_columns(headers)
            if not (columns.get("title") or columns.get("text")):
                skipped.append(f"{path.name}:{worksheet.title}:缺少标题/正文")
                continue
            label_column = choose_label_column(headers)
            for row_number, row in enumerate(worksheet.iter_rows(min_row=2), start=2):
                values = {headers[index]: row[index].value for index in range(len(headers))}
                label = label_from_color(row)
                if label is None and label_column:
                    label = label_from_value(values.get(label_column))
                if label is None:
                    skipped.append(f"{path.name}:{worksheet.title}:{row_number}:无法识别标签")
                    continue

                title = clean(values.get(columns.get("title", ""), ""))
                text = clean(values.get(columns.get("text", ""), ""))
                cover_ocr = clean(values.get(columns.get("cover_ocr", ""), ""))
                content_ocr = clean(values.get(columns.get("content_ocr", ""), ""))
                combined = "\n".join(part for part in (title, text, cover_ocr, content_ocr) if part)
                if not combined:
                    skipped.append(f"{path.name}:{worksheet.title}:{row_number}:文本为空")
                    continue

                sample_id = f"b{len(records) + 1:06d}"
                records.append(
                    {
                        "id": sample_id,
                        "title": title,
                        "text": text,
                        "cover_ocr": cover_ocr,
                        "content_ocr": content_ocr,
                        "combined_text": combined,
                        "ocr_text": "\n".join(part for part in (cover_ocr, content_ocr) if part),
                        "label": label,
                        "category": infer_category(path.name, label),
                        "brands": extract_brands(combined),
                        "note": "人工标注：来自 Excel 行填充颜色",
                        "data_scope": clean(values.get("数据范围", "")),
                    }
                )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("w", encoding="utf-8") as handle:
        for record in records:
            handle.write(json.dumps(record, ensure_ascii=False) + "\n")

    print(f"converted_rows={len(records)}")
    print(f"output={output_path.resolve()}")
    if skipped:
        print(f"skipped_rows={len(skipped)}")
        for item in skipped[:20]:
            print(f"  {item}")
        if len(skipped) > 20:
            print(f"  ... 其余 {len(skipped) - 20} 条未展开")
        # Empty or malformed source rows are reported as warnings so a normal
        # batch with a few blank rows still completes successfully.
        return 0
    return 0


def main() -> int:
    parser = argparse.ArgumentParser(description="转换人工标注 Excel Benchmark 为 JSONL")
    parser.add_argument("--input-dir", default="benchmark")
    parser.add_argument(
        "--output",
        default="data/beer_sentiment_benchmark_real.jsonl",
        help="输出 JSONL（默认写入被 Git 忽略的本地 data/ 目录）",
    )
    parser.add_argument("--inspect", action="store_true", help="只检查工作簿结构，不生成 JSONL")
    args = parser.parse_args()
    input_dir = Path(args.input_dir)
    if not input_dir.exists():
        parser.error(f"输入目录不存在：{input_dir}")
    if args.inspect:
        return inspect(input_dir)
    return convert(input_dir, Path(args.output))


if __name__ == "__main__":
    raise SystemExit(main())
